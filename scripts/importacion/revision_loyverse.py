"""Genera el Excel de revisión de datos de Loyverse para COVECA.
Uso: python3 revision_loyverse.py articulos.csv clientes.csv salida.xlsx
"""
import csv, re, sys, unicodedata
from datetime import datetime
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

art_csv, cli_csv, salida = sys.argv[1:4]
P = list(csv.DictReader(open(art_csv, encoding='utf-8')))
C = list(csv.DictReader(open(cli_csv, encoding='utf-8')))

def num(x):
    try: return float(str(x).replace(',', '.'))
    except Exception: return None

def norm(s):
    return unicodedata.normalize('NFKD', (s or '').lower()).encode('ascii', 'ignore').decode().strip()

# ---------- Correcciones de nombres (errores de tipeo evidentes) ----------
CORRECCIONES = {
    'Gillete': 'Gillette',
    'Aclado Nat': 'Aclarado Nat',
    'Mazana': 'Manzana',
    'Amasanderi ': 'Amasandería ',
}
def corregir(nombre):
    nuevo = nombre
    for a, b in CORRECCIONES.items():
        nuevo = nuevo.replace(a, b)
    return nuevo.strip()

# ---------- RUT chileno ----------
def rut_valido(r):
    s = re.sub(r'[^0-9kK]', '', r or '')
    if len(s) < 2: return None
    cuerpo, dv = s[:-1], s[-1].upper()
    suma, mult = 0, 2
    for d in reversed(cuerpo):
        suma += int(d) * mult; mult = 2 if mult == 7 else mult + 1
    calc = 11 - suma % 11
    calc = '0' if calc == 11 else 'K' if calc == 10 else str(calc)
    return (f"{int(cuerpo):,}".replace(',', '.') + '-' + dv, calc == dv)

COMUNAS = {'cauquenes': 'Cauquenes', 'retiro': 'Retiro', 'san carlos': 'San Carlos', 'parral': 'Parral',
           'pelluhue': 'Pelluhue', 'chanco': 'Chanco', 'buli': 'Buli', 'talquita': 'Talquita'}
def comuna(c):
    for campo in (c['Ciudad'], c['Dirección']):
        n = norm(campo)
        for k, v in COMUNAS.items():
            if k in n: return v
    return ''

RUTAS = ['Cauquenes', 'Retiro', 'San Carlos', 'Parral', 'Pelluhue / Chanco']
def ruta_de(com):
    if com in ('Pelluhue', 'Chanco'): return 'Pelluhue / Chanco'
    return com if com in RUTAS else ''

# ---------- Estilos ----------
F = 'Arial'
fuente = Font(name=F, size=10)
negrita = Font(name=F, size=10, bold=True)
titulo = Font(name=F, size=14, bold=True, color='3E4095')
cab = Font(name=F, size=10, bold=True, color='FFFFFF')
fill_cab = PatternFill('solid', fgColor='3E4095')
fill_input = PatternFill('solid', fgColor='FFF2A8')
fill_alerta = PatternFill('solid', fgColor='F8D7DA')
fill_aviso = PatternFill('solid', fgColor='FCE8C8')
fino = Side(style='thin', color='D9DCE5')
borde = Border(bottom=fino)
envolver = Alignment(wrap_text=True, vertical='top')

def encabezado(ws, fila, cols, anchos):
    for i, (t, w) in enumerate(zip(cols, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font, c.fill, c.alignment = cab, fill_cab, Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[fila].height = 30
    ws.freeze_panes = ws.cell(row=fila + 1, column=3)

wb = Workbook()

# =====================================================================
# Hoja 1: Cómo usar
# =====================================================================
g = wb.active
g.title = 'Cómo usar'
g.column_dimensions['A'].width = 46
g.column_dimensions['B'].width = 16
g.column_dimensions['C'].width = 70
g['A1'] = 'Revisión de datos de Loyverse antes de importar a COVECA'; g['A1'].font = titulo
g['A2'] = f'Fuente: exportaciones de Loyverse (artículos y clientes) del {datetime.now():%d-%m-%Y}.'; g['A2'].font = fuente
g['A4'] = 'Cómo completar este archivo'; g['A4'].font = negrita
pasos = [
    'Solo edita las celdas AMARILLAS. Todo lo demás se calcula solo o viene de Loyverse.',
    'Productos: revisa el nombre propuesto, el stock inicial y si mantienes el precio o lo ajustas al margen objetivo.',
    'Clientes: elige la ruta de cada cliente, cada cuántos días se visita y si se importa (Sí/No).',
    'Rutas: completa el día de la semana de cada ruta. Puedes cambiar o agregar nombres de ruta.',
    'Cuando termines, súbelo al Drive o adjúntalo en el chat. Claude importa lo aprobado.',
]
for i, p in enumerate(pasos, start=5):
    g.cell(row=i, column=1, value=f'{i-4}. {p}').font = fuente
    g.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)

g['A11'] = 'Margen objetivo (sobre precio de venta)'; g['A11'].font = negrita
g['B11'] = 0.20; g['B11'].number_format = '0%'; g['B11'].fill = fill_input; g['B11'].font = Font(name=F, size=10, color='0000FF', bold=True)
g['C11'] = 'Regla COVECA: precio = costo ÷ (1 − margen). Ej.: $1.000 ÷ 0,8 = $1.250 (20 %). Dato entregado por Jose Daniel.'
g['C11'].font = fuente

g['A13'] = 'Resumen'; g['A13'].font = negrita
filas_resumen = [
    ('Productos en Loyverse', "=COUNTA(Productos!A2:A500)", ''),
    ('Productos con stock negativo', "=COUNTIF(Productos!I2:I500,\"<0\")", 'Ventas registradas sin ingreso de mercadería. Se propone stock inicial 0 hasta contar en bodega.'),
    ('Productos bajo el margen objetivo', "=COUNTIF(Productos!H2:H500,\"<\"&$B$11)", 'Con el margen de la celda B11. Decide en la columna "Acción precio".'),
    ('Productos que se ajustarán al margen', "=COUNTIF(Productos!M2:M500,\"Ajustar a margen\")", ''),
    ('Productos sin categoría', "=COUNTIF(Productos!D2:D500,\"(sin categoría)\")", 'Asigna una categoría en la columna amarilla.'),
    ('Clientes en Loyverse', "=COUNTA(Clientes!A2:A500)", ''),
    ('Clientes que se importarán', "=COUNTIF(Clientes!N2:N500,\"Sí\")", ''),
    ('Clientes sin ruta asignada (que se importan)', "=COUNTIFS(Clientes!N2:N500,\"Sí\",Clientes!L2:L500,\"\")", 'Hay que asignarles ruta antes de importar.'),
]
for i, (t, f, nota) in enumerate(filas_resumen, start=14):
    g.cell(row=i, column=1, value=t).font = fuente
    c = g.cell(row=i, column=2, value=f); c.font = negrita
    g.cell(row=i, column=3, value=nota).font = fuente

g['A23'] = 'Colores'; g['A23'].font = negrita
g['A24'] = 'Amarillo'; g['A24'].fill = fill_input; g['B24'] = 'Para completar'; 
g['A25'] = 'Rojo'; g['A25'].fill = fill_alerta; g['B25'] = 'Problema a resolver'
g['A26'] = 'Naranja'; g['A26'].fill = fill_aviso; g['B26'] = 'Revisar'
for r in range(24, 27):
    for col in 'AB': g[f'{col}{r}'].font = fuente

# =====================================================================
# Hoja 2: Productos
# =====================================================================
ws = wb.create_sheet('Productos')
cols = ['REF', 'Nombre en Loyverse', 'Nombre propuesto', 'Categoría', 'Costo ($)', 'Precio ($)',
        'Ganancia por unidad ($)', 'Margen actual', 'Stock Loyverse', 'Stock inicial', 'Precio al margen objetivo ($)',
        'Diferencia de precio ($)', 'Acción precio', 'Precio final ($)', 'Observaciones']
encabezado(ws, 1, cols, [9, 38, 38, 16, 11, 11, 12, 10, 10, 10, 14, 12, 18, 13, 48])
cats = sorted({p['Categoria'] for p in P if p['Categoria'].strip()})
fila = 2
for p in sorted(P, key=lambda x: x['Nombre'].lower()):
    obs = []
    costo = num(p['Coste']) or 0
    precio_txt = p['Precio [COVECA]'].strip()
    precio = num(precio_txt)
    stock = num(p['En inventario [COVECA]'])
    propuesto = corregir(p['Nombre'])
    if propuesto != p['Nombre']: obs.append('Corrección de tipeo en el nombre.')
    if precio is None:
        obs.append('Precio "variable" en Loyverse y costo 0: define costo y precio.')
    if p['Seguir el Inventario'] == 'N':
        obs.append('En Loyverse no se controlaba su stock.')
    if stock is not None and stock < 0:
        obs.append('Stock negativo: se vendió más de lo registrado. Contar en bodega.')
    if 'rolls' in p['Nombre'].lower() and '150gr' in p['Handle']:
        obs.append('Revisar gramaje: el nombre dice 100gr y el código interno 150gr.')
    if precio and costo and precio < costo:
        obs.append('Se vende bajo el costo.')
    r = fila
    vals = [p['REF'], p['Nombre'], propuesto, p['Categoria'].strip() or '(sin categoría)', costo,
            precio if precio is not None else None,
            f'=IF(F{r}="","",F{r}-E{r})',
            f'=IF(OR(F{r}="",F{r}=0),"",(F{r}-E{r})/F{r})',
            stock if stock is not None else None,
            max(0, int(stock)) if stock is not None else 0,
            f"=IF(E{r}=0,\"\",ROUND(E{r}/(1-'Cómo usar'!$B$11),0))",
            f'=IF(OR(K{r}="",F{r}=""),"",K{r}-F{r})',
            'Mantener',
            f'=IF(M{r}="Ajustar a margen",K{r},F{r})',
            ' '.join(obs)]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=ci, value=v)
        c.font = fuente; c.border = borde
    for ci in (5, 6, 7, 11, 12, 14):
        ws.cell(row=r, column=ci).number_format = '$#,##0;-$#,##0;-'
    ws.cell(row=r, column=8).number_format = '0.0%'
    for ci in (9, 10):
        ws.cell(row=r, column=ci).number_format = '#,##0;-#,##0;0'
    for ci in (3, 4, 10, 13):
        ws.cell(row=r, column=ci).fill = fill_input
    if precio is None:
        for ci in (5, 6): ws.cell(row=r, column=ci).fill = fill_input
    ws.cell(row=r, column=15).alignment = envolver
    fila += 1
ultp = fila - 1
dv_acc = DataValidation(type='list', formula1='"Mantener,Ajustar a margen"', allow_blank=False)
dv_cat = DataValidation(type='list', formula1='"' + ','.join(cats) + '"', allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dv_acc); ws.add_data_validation(dv_cat)
dv_acc.add(f"M2:M{ultp}"); dv_cat.add(f"D2:D{ultp}")
ws.conditional_formatting.add(f'I2:I{ultp}', CellIsRule(operator='lessThan', formula=['0'], fill=fill_alerta))
ws.conditional_formatting.add(f'H2:H{ultp}', FormulaRule(formula=[f"AND(H2<>\"\",H2<0)"], fill=fill_alerta))
ws.conditional_formatting.add(f'H2:H{ultp}', FormulaRule(formula=[f"AND(H2<>\"\",H2<'Cómo usar'!$B$11)"], fill=fill_aviso))
ws.conditional_formatting.add(f'D2:D{ultp}', CellIsRule(operator='equal', formula=['"(sin categoría)"'], fill=fill_alerta))
ws.auto_filter.ref = f'A1:O{ultp}'
ws['C1'].comment = Comment('Corrige si quieres otro nombre. Así aparecerá en la app y en la nota.', 'Claude')
ws['J1'].comment = Comment('Stock con el que parte la app. Negativos se proponen en 0: ideal hacer un conteo en bodega antes de partir.', 'Claude')
ws['M1'].comment = Comment('Mantener = sigue el precio de Loyverse. Ajustar a margen = usa el precio de la columna K.', 'Claude')

# =====================================================================
# Hoja 3: Rutas
# =====================================================================
wr = wb.create_sheet('Rutas')
encabezado(wr, 1, ['Ruta', 'Día de visita', 'Clientes asignados', 'Notas'], [24, 16, 18, 60])
wr.freeze_panes = 'A2'
dv_dia = DataValidation(type='list', formula1='"Lunes,Martes,Miércoles,Jueves,Viernes,Sábado"', allow_blank=True)
wr.add_data_validation(dv_dia)
notas_ruta = {'Pelluhue / Chanco': 'Agrupa Pelluhue y Chanco (pocos clientes). Sepáralas si son días distintos.'}
for i, rt in enumerate(RUTAS + ['', '', ''], start=2):
    wr.cell(row=i, column=1, value=rt or None).fill = fill_input
    wr.cell(row=i, column=2).fill = fill_input
    wr.cell(row=i, column=3, value=f'=IF(A{i}="","",COUNTIFS(Clientes!L$2:L$500,A{i},Clientes!N$2:N$500,"Sí"))')
    wr.cell(row=i, column=4, value=notas_ruta.get(rt))
    for ci in range(1, 5): wr.cell(row=i, column=ci).font = fuente
dv_dia.add('B2:B10')
wr['A12'] = 'Puedes renombrar rutas o agregar nuevas en las filas vacías; la lista de rutas de la hoja Clientes se actualiza sola.'
wr['A12'].font = Font(name=F, size=10, italic=True)

# =====================================================================
# Hoja 4: Clientes
# =====================================================================
wc = wb.create_sheet('Clientes')
cols = ['ID Loyverse', 'Nombre en Loyverse', 'Nombre propuesto', 'Correo', 'Teléfono', 'Dirección', 'Comuna detectada',
        'RUT', 'Última compra', 'Compras', 'Total comprado ($)', 'Ruta', 'Frecuencia (días)', 'Importar', 'Observaciones']
encabezado(wc, 1, cols, [11, 34, 34, 28, 13, 26, 15, 14, 13, 9, 15, 18, 12, 10, 50])
dups = {n for n, k in Counter(norm(c['Nombre del cliente']) for c in C).items() if k > 1}
hace6 = datetime(2026, 4, 9)
fila = 2
for c in sorted(C, key=lambda x: -float(x['Gasto total'] or 0)):
    obs = []
    nombre = c['Nombre del cliente'].strip()
    propuesto = corregir(nombre + ' ').strip()
    if propuesto != nombre: obs.append('Corrección de tipeo en el nombre.')
    if norm(nombre) in dups: obs.append('Nombre repetido en Loyverse: revisar si es la misma persona.')
    rut_txt = ''
    if c['Código de cliente'].strip():
        rv = rut_valido(c['Código de cliente'])
        if rv and rv[1]: rut_txt = rv[0]
        else: obs.append(f'Código "{c["Código de cliente"]}" no es un RUT válido.')
    correo = c['Email'].strip()
    if correo.endswith('.ccom'):
        obs.append('Correo con error (.ccom): se corrige a .com.'); correo = correo[:-5] + '.com'
    ult_txt = c['Ultima visita'].split(' ')[0]
    try: ult = datetime.strptime(ult_txt, '%d-%m-%y')
    except Exception: ult = None
    importar = 'Sí'
    if ult and ult < hace6:
        obs.append('Sin compras hace más de 6 meses.')
    com = comuna(c)
    if not com: obs.append('Sin ciudad en Loyverse: asignar ruta.')
    if 'el jaque' in norm(nombre): obs.append('Mismo negocio con dos sucursales: se mantienen separadas.')
    r = fila
    direccion = c['Dirección'].strip()
    if norm(direccion) in COMUNAS: direccion = ''
    vals = [int(c['ID del cliente']), nombre, propuesto, correo or None, c['Número de teléfono'].strip() or None,
            direccion or None, com or None, rut_txt or None, ult, int(c['Total de visitas'] or 0),
            int(float(c['Gasto total'] or 0)), ruta_de(com) or None, 7, importar, ' '.join(obs) or None]
    for ci, v in enumerate(vals, start=1):
        cell = wc.cell(row=r, column=ci, value=v); cell.font = fuente; cell.border = borde
    wc.cell(row=r, column=9).number_format = 'dd-mm-yyyy'
    wc.cell(row=r, column=11).number_format = '$#,##0'
    for ci in (3, 5, 8, 12, 13, 14):
        wc.cell(row=r, column=ci).fill = fill_input
    wc.cell(row=r, column=15).alignment = envolver
    fila += 1
ultc = fila - 1
dv_ruta = DataValidation(type='list', formula1='=Rutas!$A$2:$A$10', allow_blank=True)
dv_frec = DataValidation(type='list', formula1='"7,14,21,30"', allow_blank=False, showErrorMessage=False)
dv_imp = DataValidation(type='list', formula1='"Sí,No"', allow_blank=False)
for dv, rng in ((dv_ruta, f'L2:L{ultc}'), (dv_frec, f'M2:M{ultc}'), (dv_imp, f'N2:N{ultc}')):
    wc.add_data_validation(dv); dv.add(rng)
wc.conditional_formatting.add(f'L2:L{ultc}', FormulaRule(formula=['AND(L2="",N2="Sí")'], fill=fill_alerta))
wc.conditional_formatting.add(f'I2:I{ultc}', CellIsRule(operator='lessThan', formula=['DATE(2026,4,9)'], fill=fill_aviso))
wc.auto_filter.ref = f'A1:O{ultc}'
wc['L1'].comment = Comment('Elige la ruta de la lista (se toma de la hoja Rutas). Rojo = falta asignar.', 'Claude')
wc['M1'].comment = Comment('Cada cuántos días se visita: 7 = semanal, 14 = quincenal.', 'Claude')
wc['N1'].comment = Comment('No = el cliente no pasa a la app nueva (su historial igual queda guardado).', 'Claude')

wb.move_sheet('Rutas', offset=1)
wb.save(salida)
print('ok', ultp - 1, 'productos', ultc - 1, 'clientes')
